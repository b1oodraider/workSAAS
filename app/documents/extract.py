"""Resume file -> plain text. Supports PDF, DOCX, TXT/MD."""

from __future__ import annotations

import io
import zipfile
from pathlib import PurePath


MAX_PDF_PAGES = 15
MAX_DOCX_UNPACKED = 50 * 1024 * 1024


class ExtractError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext in {".txt", ".md", ""}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise ExtractError(f"Формат {ext} не поддерживается: загрузите PDF, DOCX или TXT")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < 50:
        raise ExtractError(
            "Не удалось извлечь текст (скан без текстового слоя?). Вставьте текст резюме вручную."
        )
    return text


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ExtractError(f"В PDF больше {MAX_PDF_PAGES} страниц — это точно резюме?")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 - pypdf raises many exception types
        raise ExtractError(f"Не удалось прочитать PDF: {exc}") from exc


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            unpacked = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise ExtractError("Не удалось прочитать DOCX: файл повреждён") from exc
    if unpacked > MAX_DOCX_UNPACKED:  # zip bomb: python-docx would inflate it all into memory
        raise ExtractError("DOCX слишком большой после распаковки — сохраните резюме в PDF или TXT")
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"Не удалось прочитать DOCX: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
